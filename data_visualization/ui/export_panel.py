"""Download PNG/SVG/HTML/PDF/Python/Notebook/PPTX."""

from __future__ import annotations

import io
import json
from typing import Optional

import pandas as pd
import plotly.graph_objects as go
import plotly.io as pio
import streamlit as st


@st.cache_data(show_spinner=False)
def _render_figure_to_bytes(fig_json: str, fmt: str, width: int = 1200, height: int = 800) -> bytes:
    """Convert a Plotly figure to image bytes — cached so Kaleido only runs when the
    figure actually changes, not on every Streamlit rerun.

    ``fig_json`` is the JSON serialization of the figure (used as the cache key).
    We serialize+deserialize so the cache is keyed purely on figure content,
    not on the in-memory object identity.
    """
    fig = pio.from_json(fig_json)
    return fig.to_image(format=fmt, width=width, height=height)


def render_export_panel(
    fig: Optional[go.Figure],
    chart_mode: str,
    chart_type: str,
    selected_table: str,
    x_col: Optional[str],
    y_col: Optional[str],
    color_col: Optional[str],
) -> None:
    if fig is None:
        return

    st.divider()
    st.subheader("📥 Export chart")
    st.caption("Download the current chart in multiple formats.")
    export_chart_name = chart_mode if chart_mode != "basic" else chart_type
    export_height = 800

    # Serialize once for all cached image calls below.
    try:
        fig_json = fig.to_json()
    except Exception as e:
        st.error(f"Could not serialize figure for export: {e}")
        return

    col1, col2, col3 = st.columns(3)
    with col1:
        try:
            img_bytes = _render_figure_to_bytes(fig_json, "png", 1200, export_height)
            st.download_button(
                "📥 Download PNG",
                img_bytes,
                f"chart_{export_chart_name}_{selected_table}.png",
                "image/png",
                key="download_png",
                width="stretch",
            )
        except Exception as e:
            st.error(f"PNG export failed: {e}")
            st.caption("Install kaleido: `pip install kaleido`")

    with col2:
        try:
            svg_bytes = _render_figure_to_bytes(fig_json, "svg", 1200, export_height)
            st.download_button(
                "📐 Download SVG",
                svg_bytes,
                f"chart_{export_chart_name}_{selected_table}.svg",
                "image/svg+xml",
                key="download_svg",
                width="stretch",
            )
        except Exception as e:
            st.error(f"SVG export failed: {e}")

    with col3:
        try:
            html_str = fig.to_html(full_html=False)
            st.download_button(
                "🌐 Download HTML",
                html_str.encode(),
                f"chart_{export_chart_name}_{selected_table}.html",
                "text/html",
                key="download_html",
                width="stretch",
            )
        except Exception as e:
            st.error(f"HTML export failed: {e}")

    st.markdown("---")
    st.subheader("📦 More formats")
    e1, e2, e3 = st.columns(3)
    with e1:
        try:
            pdf_bytes = _render_figure_to_bytes(fig_json, "pdf", 1200, export_height)
            st.download_button(
                "📄 Download PDF",
                pdf_bytes,
                f"chart_{export_chart_name}_{selected_table}.pdf",
                "application/pdf",
                key="download_pdf",
                width="stretch",
            )
        except Exception as e:
            st.error(f"PDF export failed: {e}")

    xr = repr(x_col)
    yr = repr(y_col)
    cr = repr(color_col)
    with e2:
        python_script = f"""import pandas as pd
import plotly.express as px

# df = pd.read_csv("{selected_table}_current.csv")

fig = px.{chart_type}(
    df,
    x={xr},
    y={yr},
    color={cr},
)
fig.show()
"""
        st.download_button(
            "🐍 Download Python",
            python_script.encode(),
            f"chart_{export_chart_name}_{selected_table}.py",
            "text/x-python",
            key="download_python",
            width="stretch",
        )

    with e3:
        notebook = {
            "cells": [
                {"cell_type": "markdown", "metadata": {}, "source": [f"# Chart: {export_chart_name}\n"]},
                {
                    "cell_type": "code",
                    "execution_count": None,
                    "metadata": {},
                    "outputs": [],
                    "source": ["import pandas as pd\n", "import plotly.express as px\n"],
                },
                {
                    "cell_type": "code",
                    "execution_count": None,
                    "metadata": {},
                    "outputs": [],
                    "source": [
                        f"fig = px.{chart_type}(df, x={xr}, y={yr}, color={cr})\n",
                        "fig.show()\n",
                    ],
                },
            ],
            "metadata": {
                "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
                "language_info": {"name": "python", "version": "3.x"},
            },
            "nbformat": 4,
            "nbformat_minor": 5,
        }
        st.download_button(
            "📓 Download Notebook",
            json.dumps(notebook, indent=2).encode(),
            f"chart_{export_chart_name}_{selected_table}.ipynb",
            "application/json",
            key="download_notebook",
            width="stretch",
        )

    ppt_col, _, _ = st.columns(3)
    with ppt_col:
        try:
            from pptx import Presentation

            img_bytes = _render_figure_to_bytes(fig_json, "png", 1200, export_height)
            prs = Presentation()
            slide_layout = prs.slide_layouts[5]
            slide = prs.slides.add_slide(slide_layout)
            image_stream = io.BytesIO(img_bytes)
            slide.shapes.add_picture(image_stream, left=0, top=0, width=prs.slide_width)
            pptx_stream = io.BytesIO()
            prs.save(pptx_stream)
            st.download_button(
                "📊 Download PowerPoint",
                pptx_stream.getvalue(),
                f"chart_{export_chart_name}_{selected_table}.pptx",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                key="download_pptx",
                width="stretch",
            )
        except ImportError:
            st.caption("Install python-pptx: `pip install python-pptx`")
        except Exception as e:
            st.error(f"PowerPoint export failed: {e}")
